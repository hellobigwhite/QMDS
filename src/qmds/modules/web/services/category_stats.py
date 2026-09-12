"""网站分类统计服务 — 主数据文件夹全表分类汇总

由数据分配任务（data_allocator.run_allocation_task）内嵌调用：分配（含拆表）
完成后，对每个网站数据文件夹（每个主分类的文件夹，含主数据表 main*.xlsx
与补充数据表 *_supp*.xlsx）自动执行统计；site_info_generator 也会在缺少
统计表时用本模块补生成。

流程：
1. 递归扫描文件夹中所有 .xlsx 数据表格，流式统计分类列（Categories/分类）
   中每个分类的产品数量（跨表累加）；
2. 汇总结果写入该文件夹下的 分类统计.xlsx：
   - Sheet「分类统计」：分类（原始值，||| 分隔层级）、一/二/三级分类、
     产品数、占比，按产品数降序 —— 供后续模型读取网站分类结构；
   - Sheet「汇总」：文件夹、统计时间、表格数、产品总数、分类总数等概要。

写入的 分类统计.xlsx 由 site_info_generator 读取，用于 AI 生成网站信息。
"""

import os
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from qmds.utils import winpath
from qmds.utils.logger import get_logger

log = get_logger("web.category_stats")

# 统计结果文件名（扫描时自动跳过，避免把统计表当成数据表重复统计）
STATS_FILE_NAME = "分类统计.xlsx"

# 分类层级分隔符（导出表格中多级分类的原始格式，如 Hardware|||Plumbing）
CATEGORY_SEP = "|||"

# 展开的分类层级列（更深的层级并入最后一级）
LEVEL_COLUMNS = ("一级分类", "二级分类", "三级分类")

# 分类统计表的数据列
STATS_COLUMNS = ("分类", *LEVEL_COLUMNS, "产品数", "占比(%)")

# AI 生成网站信息写入的结果表格（批量任务汇总表/单网站表，扫描时一并跳过）
INFO_FILE_NAME = "网站信息.xlsx"


def split_category_levels(value) -> list[str]:
    """拆分分类层级：'Hardware|||Plumbing & Fittings' -> ['Hardware', 'Plumbing & Fittings']

    超过 3 级时，更深的层级以 ' > ' 连接并入第三级，
    保证任意分类都能完整还原。空值返回 []。
    """
    sval = str(value or "").strip()
    if not sval:
        return []
    parts = [p.strip() for p in sval.split(CATEGORY_SEP) if p.strip()]
    if not parts:
        return []
    if len(parts) > len(LEVEL_COLUMNS):
        head = parts[:len(LEVEL_COLUMNS) - 1]
        parts = head + [" > ".join(parts[len(LEVEL_COLUMNS) - 1:])]
    return parts


def _is_scannable(path: Path) -> bool:
    """可统计的表格：.xlsx、非 Excel 临时锁文件、非结果文件（统计表/网站信息表）"""
    return (path.suffix.lower() == ".xlsx"
            and not path.name.startswith("~$")
            and path.name not in (STATS_FILE_NAME, INFO_FILE_NAME))


def _rglob_xlsx(folder: Path) -> list[Path]:
    """递归列出文件夹下所有 .xlsx（超长路径安全）

    os.scandir 在 Windows 上支持超过 260 字符的目录路径（walk 用的就是它），
    而 Path.rglob 在部分 Python 版本上对超长路径会抛 OSError 或漏文件。
    数据分配的输出（分类文件夹/分卷文件）很容易超过 260 字符。
    """
    folder = Path(folder)
    top = winpath.long_path(folder)
    out: list[Path] = []
    for root, _dirs, names in os.walk(top):
        # relpath 还原相对/绝对形式（保持入参路径形态，不引入前缀）
        rel = os.path.relpath(root, top)
        base = folder if rel == "." else folder / Path(rel)
        for n in names:
            if n.lower().endswith(".xlsx"):
                out.append(base / n)
    return out


def collect_stats_files(folder) -> list[Path]:
    """收集文件夹（含子文件夹）中所有可统计的 .xlsx，按自然顺序排列"""
    folder = Path(folder)
    if not winpath.is_dir(folder):
        return []
    files = [p for p in _rglob_xlsx(folder) if _is_scannable(p)]
    try:
        files.sort(key=lambda p: [int(t) if t.isdigit() else t.lower()
                                  for t in re.split(r"(\d+)", str(p))])
    except TypeError:
        # 数字/文本 token 同位比较冲突时回退字典序
        files.sort(key=lambda p: str(p).lower())
    return files


def count_file_categories(filepath: Path):
    """流式统计单个表格的分类列

    返回 (分类计数 Counter, 总行数, 空分类行数)；
    无分类列时抛 ValueError（调用方决定跳过）。
    """
    from openpyxl import load_workbook

    from qmds.modules.web.services.data_allocator import detect_category_column

    wb = load_workbook(winpath.long_path(filepath), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("表格为空或格式不正确")

        col = detect_category_column(header)
        if col is None:
            raise ValueError("表格中未找到分类列（Categories/分类）")

        idx = list(header).index(col)
        counter: Counter = Counter()
        total = 0
        empty = 0
        for row in rows:
            if row is None:
                continue
            if all(v is None or str(v).strip() == "" for v in row):
                continue  # 跳过整行空行
            total += 1
            val = row[idx] if idx < len(row) else None
            sval = str(val).strip() if val is not None else ""
            if sval:
                counter[sval] += 1
            else:
                empty += 1
        return counter, total, empty
    finally:
        wb.close()


def aggregate_folder_categories(files, log_fn=None, stop_check=None) -> dict:
    """聚合多个表格的分类统计

    返回:
        {
            "counts": Counter{分类: 产品数}（跨表累加，不含空分类）,
            "total_rows": 数据总行数,
            "empty_rows": 空分类行数,
            "files": [(文件名, 行数), ...] 按输入顺序,
            "skipped": [(文件名, 原因), ...],
        }
    """
    counts: Counter = Counter()
    total_rows = 0
    empty_rows = 0
    file_rows: list[tuple[str, int]] = []
    skipped: list[tuple[str, str]] = []

    for fp in files:
        if stop_check and stop_check():
            raise InterruptedError("任务被用户停止")
        try:
            counter, total, empty = count_file_categories(Path(fp))
        except Exception as e:  # 单个表格失败不影响整体
            skipped.append((Path(fp).name, str(e)))
            if log_fn:
                log_fn(f"跳过表格 {Path(fp).name}: {e}", "warning")
            continue
        counts.update(counter)
        total_rows += total
        empty_rows += empty
        file_rows.append((Path(fp).name, total))
        if log_fn:
            log_fn(f"已统计 {Path(fp).name}: {total} 行，{len(counter)} 个分类")

    return {
        "counts": counts,
        "total_rows": total_rows,
        "empty_rows": empty_rows,
        "files": file_rows,
        "skipped": skipped,
    }


def build_stats_rows(counts: dict) -> list[dict]:
    """分类计数 -> 统计表行（按产品数降序，含层级拆分与占比）"""
    classified = sum(counts.values())
    rows = []
    for cat, cnt in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
        levels = split_category_levels(cat)
        row = {
            "分类": cat,
            "产品数": int(cnt),
            "占比(%)": round(cnt / classified * 100, 2) if classified else 0.0,
        }
        for i, col in enumerate(LEVEL_COLUMNS):
            row[col] = levels[i] if i < len(levels) else ""
        rows.append(row)
    return rows


def write_stats_excel(out_path: Path, agg: dict, folder_label: str = "") -> Path:
    """把聚合结果写入 分类统计.xlsx（两个 Sheet：分类统计 + 汇总）"""
    import pandas as pd

    out_path = Path(out_path)
    counts = agg["counts"]
    rows = build_stats_rows(counts)
    classified = sum(counts.values())
    level1 = {split_category_levels(c)[0] for c in counts if split_category_levels(c)}

    df_stats = pd.DataFrame(rows, columns=list(STATS_COLUMNS))
    df_summary = pd.DataFrame([
        ("数据文件夹", folder_label or out_path.parent.name),
        ("统计时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("数据表格数", len(agg.get("files", []))),
        ("跳过表格数", len(agg.get("skipped", []))),
        ("产品总数", agg.get("total_rows", 0)),
        ("有分类产品数", classified),
        ("未分类产品数", agg.get("empty_rows", 0)),
        ("分类总数", len(counts)),
        ("一级分类数", len(level1)),
    ], columns=["指标", "值"])

    # 输出路径可能超过 Windows 260 字符限制（分类名很长时），用扩展前缀写出
    with pd.ExcelWriter(winpath.long_path(out_path), engine="openpyxl") as writer:
        df_stats.to_excel(writer, sheet_name="分类统计", index=False)
        df_summary.to_excel(writer, sheet_name="汇总", index=False)
        ws = writer.book["分类统计"]
        for i, col in enumerate(STATS_COLUMNS, start=1):
            width = max(len(col) * 2 + 2, 14)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    return out_path


def _safe_int(v, default: int = 0) -> int:
    """安全转 int（None/NaN/非法值返回 default）"""
    try:
        if v is None:
            return default
        f = float(v)
        if f != f:  # NaN
            return default
        return int(f)
    except (TypeError, ValueError):
        return default


def read_stats_excel(path) -> dict:
    """读取 分类统计.xlsx，还原网站分类结构（供模型/后续功能使用）

    返回:
        {
            "categories": [{"category", "count", "level1", "level2", "level3"}, ...],
            "summary": {指标: 值},
        }
    """
    import pandas as pd

    path = Path(path)
    if not winpath.is_file(path):
        raise FileNotFoundError(f"分类统计表不存在: {path.name}")

    df = pd.read_excel(winpath.long_path(path), sheet_name="分类统计", engine="openpyxl")
    categories = []
    for _, row in df.iterrows():
        cat = str(row.get("分类") or "").strip()
        if not cat or cat.lower() == "nan":
            continue
        levels = split_category_levels(cat)  # 优先用原始分类列还原层级
        if not levels:
            levels = [str(row.get(c) or "").strip() for c in LEVEL_COLUMNS]
            levels = [v for v in levels if v and v.lower() != "nan"]
        categories.append({
            "category": cat,
            "count": _safe_int(row.get("产品数")),
            "level1": levels[0] if len(levels) > 0 else "",
            "level2": levels[1] if len(levels) > 1 else "",
            "level3": levels[2] if len(levels) > 2 else "",
        })

    summary: dict = {}
    try:
        df_sum = pd.read_excel(winpath.long_path(path), sheet_name="汇总", engine="openpyxl")
        for _, row in df_sum.iterrows():
            key = str(row.get("指标") or "").strip()
            if key:
                summary[key] = row.get("值")
    except Exception:
        pass

    return {"categories": categories, "summary": summary}
