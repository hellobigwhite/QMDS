"""批量拆表服务 — 移植自 BB_Data_Tool 的批量拆表工具（DataHandle/splitter_main）

功能：
- 将 Excel 按行数拆分为多个文件（rows_per_file=None 表示不限制，整表一份）；
- 输出到源文件同目录的 {源文件名}_split 文件夹；
- 文件名: {源文件名}_part{N}_{两位随机大写字母}{时间戳}.xlsx（保证唯一）；
- 可选给"原站域名"列添加后缀：
  * none   不处理
  * custom 添加自定义后缀（值_后缀）
  * part   按分卷添加 _part1 _part2 ...
- 可选拆分完成后删除源文件；
- 可选使用 xlwings（需本机安装 Excel）重新保存拆分文件，以兼容
  站群系统等旧版 PHP Excel 读取器（xlwings 不可用时自动跳过并告警）。
"""

import random
import string
import time
from pathlib import Path

from qmds.utils import winpath
from qmds.utils.logger import get_logger

log = get_logger("web.excel_splitter")

# 后缀模式
SUFFIX_NONE = "none"
SUFFIX_CUSTOM = "custom"
SUFFIX_PART = "part"
SUFFIX_MODES = (SUFFIX_NONE, SUFFIX_CUSTOM, SUFFIX_PART)

# 域名列名（与导出列一致）
DOMAIN_COLUMN = "原站域名"


def _unique_part_path(output_folder: Path, base_name: str, part_index: int) -> Path:
    """生成唯一的分卷文件路径: {base}_part{N}_{2位随机字母}{时间戳}.xlsx"""
    letters = "".join(random.choices(string.ascii_uppercase, k=2))
    timestamp = int(time.time())
    return output_folder / f"{base_name}_part{part_index}_{letters}{timestamp}.xlsx"


def _apply_suffix(header, row, domain_idx, suffix_mode, custom_suffix, part_index):
    """给一行的"原站域名"列添加后缀（值为空则跳过）"""
    if domain_idx is None or domain_idx >= len(row):
        return row
    value = row[domain_idx]
    if value is None or str(value).strip() == "":
        return row
    if suffix_mode == SUFFIX_CUSTOM and custom_suffix:
        new_value = f"{value}_{custom_suffix}"
    elif suffix_mode == SUFFIX_PART and part_index is not None:
        new_value = f"{value}_part{part_index}"
    else:
        return row
    row = list(row)
    row[domain_idx] = new_value
    return row


def split_excel_file(file_path, rows_per_file=None, suffix_mode=SUFFIX_NONE,
                     custom_suffix=None, remove_source=False,
                     resave_with_excel=False, progress_callback=None,
                     stop_check=None, expected_rows=None,
                     output_folder=None) -> dict:
    """将单个 Excel 文件按行数拆分

    参数:
        file_path: 源文件路径
        rows_per_file: 每份最大数据行数；None 或 <=0 表示不限制（整表一份）
        suffix_mode: 原站域名列后缀模式（none/custom/part）
        custom_suffix: 自定义后缀内容（suffix_mode=custom 时生效）
        remove_source: 拆分成功后删除源文件
        resave_with_excel: 使用 xlwings 重存拆分文件（需本机 Excel，兼容站群系统）
        progress_callback: 进度回调（接受该文件 0~100 int）
        stop_check: 停止检查函数（返回 True 时中止任务）
        expected_rows: 预期数据行数（用于计算该文件的进度百分比，可省略）
        output_folder: 拆分输出文件夹（默认为源文件旁的 {源文件名}_split；
                       指定时拆分结果与源文件同文件夹，用于主数据与补充数据共用目录）

    返回:
        {"output_folder": Path, "parts": [Path...], "rows": 总数据行数}
    """
    from openpyxl import Workbook, load_workbook

    if suffix_mode not in SUFFIX_MODES:
        suffix_mode = SUFFIX_NONE

    file_path = Path(file_path)
    base_name = file_path.stem
    if output_folder is None:
        output_folder = file_path.parent / f"{base_name}_split"
    output_folder = Path(output_folder)
    # 数据分配输出的分类文件夹/分卷文件路径可能超过 Windows MAX_PATH(260)，
    # 统一用 \\?\\ 扩展长度前缀创建目录（winpath.makedirs）
    winpath.makedirs(output_folder)

    unlimited = not rows_per_file or rows_per_file <= 0

    src_wb = load_workbook(winpath.long_path(file_path), read_only=True, data_only=True)
    parts = []
    total_rows = 0
    try:
        sheet = src_wb.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise ValueError(f"表格为空或格式不正确: {file_path.name}")

        # 定位原站域名列（无该列时后缀功能跳过）
        header_list = [str(c) if c is not None else "" for c in header]
        domain_idx = header_list.index(DOMAIN_COLUMN) if DOMAIN_COLUMN in header_list else None

        def save_part(buf, part_index):
            wb = Workbook()
            ws = wb.active
            ws.append(list(header))
            for row in buf:
                ws.append(_apply_suffix(header_list, row, domain_idx,
                                        suffix_mode, custom_suffix, part_index))
            out = _unique_part_path(output_folder, base_name, part_index)
            # 分卷文件名比源文件长（_partN_XX时间戳.xlsx），路径超 260 时
            # 必须用扩展长度前缀保存（winpath.long_path）；parts 列表仍存
            # 普通路径供日志与后续流程使用
            wb.save(winpath.long_path(out))
            wb.close()
            parts.append(out)
            log.info(f"已生成分卷: {out.name}（{len(buf)} 条）")

        def report_progress():
            if not progress_callback:
                return
            if expected_rows:
                progress_callback(min(99, int(total_rows / expected_rows * 100)))
            else:
                progress_callback(50)

        row_buffer = []
        part_index = 1
        for row in rows_iter:
            if row is None:
                continue
            if all(v is None or str(v).strip() == "" for v in row):
                continue  # 跳过整行空行
            row_buffer.append(row)
            total_rows += 1
            if not unlimited and len(row_buffer) >= rows_per_file:
                if stop_check and stop_check():
                    raise InterruptedError("任务被用户停止")
                save_part(row_buffer, part_index)
                row_buffer = []
                part_index += 1
                report_progress()

        if stop_check and stop_check():
            raise InterruptedError("任务被用户停止")
        if row_buffer or not parts:
            save_part(row_buffer, part_index)
    finally:
        src_wb.close()

    # 可选：使用 xlwings（本机 Excel）重存，兼容旧版 PHP Excel 读取器
    if resave_with_excel:
        resave_folder_with_excel(output_folder)

    if remove_source:
        try:
            winpath.remove(file_path)
            log.info(f"已删除源文件: {file_path.name}")
        except OSError as e:
            log.warning(f"删除源文件失败 {file_path.name}: {e}")

    if progress_callback:
        progress_callback(100)

    return {"output_folder": output_folder, "parts": parts, "rows": total_rows}


def resave_folder_with_excel(folder_path) -> int:
    """使用本机 Excel（xlwings）重新保存文件夹中的 .xlsx 文件

    移植自 BB 工具 save_with_xlwings：解决部分系统（如站群 PHP 读取器）
    对 openpyxl 生成文件兼容性差的问题。xlwings 或 Excel 不可用时跳过。
    返回成功重存的文件数。
    """
    folder = Path(folder_path)
    if not folder.is_dir():
        return 0
    try:
        import xlwings as xw
    except ImportError:
        log.warning("xlwings 未安装，跳过 Excel 兼容重存（pip install xlwings 后可用）")
        return 0

    count = 0
    app = None
    try:
        app = xw.App(visible=False)
        app.display_alerts = False
        for fp in sorted(folder.glob("*.xlsx")):
            if fp.name.startswith("~$"):
                continue
            try:
                wb = app.books.open(str(fp))
                wb.save(str(fp))
                wb.close()
                count += 1
                log.info(f"Excel 兼容重存完成: {fp.name}")
            except Exception as e:
                log.warning(f"Excel 兼容重存失败 {fp.name}: {e}")
    except Exception as e:
        log.warning(f"启动 Excel 失败，跳过兼容重存: {e}")
    finally:
        if app is not None:
            try:
                app.quit()
            except Exception:
                pass
    return count
